"""Private JSON/XLSX exports frozen at command acceptance."""

import json
from copy import deepcopy
from datetime import UTC, datetime
from io import BytesIO
from typing import Any
from urllib.parse import quote
from xml.etree import ElementTree
from zipfile import ZipFile

from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.orm import Session as DBSession
from temporalio import activity

from app import wire as w
from app.auth import Actor, reject
from app.commands import audit, command, emit, owned_run
from app.config import Settings
from app.control import accepted, proposal_row, resource_row
from app.db import (
    Audit,
    EvidenceArtifact,
    Outbox,
    Run,
    RunEvidence,
    Session,
    SourceVersion,
    StageOutput,
)
from app.storage import storage


def render_snapshot(payload: dict[str, Any], run_id: str) -> dict[str, Any]:
    result = deepcopy(payload["result"])
    render_id = payload["manifest"]["render_request_id"]
    for asset in result["assets"]:
        asset["url"] = (
            f"/api/v1/runs/{run_id}/renders/{render_id}/assets/{quote(asset['asset_id'], safe='')}"
        )
    return {"manifest": deepcopy(payload["manifest"]), "result": result}


def create_export(
    actor: Actor, run_id: str, body: w.ExportCommand, key: str
) -> w.Accepted:
    def effect(session: DBSession, command_id: str) -> dict[str, Any]:
        run = owned_run(session, actor, run_id, lock=True)
        row = proposal_row(session, run, body.proposal_set_version)
        proposals = w.Proposals.model_validate(row.payload)
        variants = {variant.variant_id: variant for variant in proposals.variants}
        if len(set(body.variant_ids)) != len(body.variant_ids) or any(
            value not in variants for value in body.variant_ids
        ):
            reject(
                "invalid_variant",
                "Export variants must belong to the selected version",
                422,
            )
        cutoff = datetime.now(UTC)
        selected = {
            change_id
            for value in body.variant_ids
            for change_id in variants[value].change_ids
        }
        frozen_proposals = proposals.model_copy(deep=True)
        frozen_proposals.variants = [
            variant
            for variant in proposals.variants
            if variant.variant_id in body.variant_ids
        ]
        frozen_proposals.changes = [
            change for change in proposals.changes if change.change_id in selected
        ]
        outputs = session.scalars(
            select(StageOutput).where(StageOutput.run_id == run_id)
        ).all()
        decisions = [
            deepcopy(value.payload)
            for value in outputs
            if value.stage_key.startswith("decision:")
            and value.payload.get("proposal_set_version") == body.proposal_set_version
        ]
        for decision in sorted(decisions, key=lambda value: str(value["occurred_at"])):
            for change in frozen_proposals.changes:
                if change.change_id in decision["change_ids"]:
                    change.decision_state = decision["state"]
        history = session.scalars(
            select(Audit)
            .where(Audit.run_id == run_id, Audit.occurred_at <= cutoff)
            .order_by(Audit.occurred_at, Audit.id)
        ).all()
        source_versions = session.scalars(
            select(SourceVersion).where(
                SourceVersion.id.in_(proposals.source_version_ids)
            )
        ).all()
        evidence_ids = {
            value
            for change in frozen_proposals.changes
            for value in change.supporting_evidence_ids + change.opposing_evidence_ids
        }
        analysis = {
            value.stage_key: deepcopy(value.payload)
            for value in outputs
            if value.generation_id == row.generation_id
            and value.stage_key
            in {
                "findings",
                "baselines",
                "research",
                "research_sources",
                "conflicts",
                "configuration",
            }
        }
        pending: list[Any] = [analysis]
        while pending:
            value = pending.pop()
            if isinstance(value, dict):
                for field, item in value.items():
                    if field == "evidence_id" and isinstance(item, str):
                        evidence_ids.add(item)
                    elif field.endswith("evidence_ids") and isinstance(item, list):
                        evidence_ids.update(item)
                    else:
                        pending.append(item)
            elif isinstance(value, list):
                pending.extend(value)
        evidence = session.scalars(
            select(EvidenceArtifact)
            .join(RunEvidence, RunEvidence.evidence_id == EvidenceArtifact.id)
            .where(
                RunEvidence.run_id == run_id,
                RunEvidence.generation_id == row.generation_id,
                EvidenceArtifact.id.in_(evidence_ids),
                EvidenceArtifact.source_version_id.in_(proposals.source_version_ids),
            )
        ).all()
        external: dict[str, w.Evidence] = {}
        for output in outputs:
            if output.generation_id != row.generation_id or not (
                output.stage_key == "research_sources"
                or output.stage_key.startswith("research:step:")
            ):
                continue
            for source in output.payload.get("sources", []):
                item = w.Evidence.model_validate(source["evidence"])
                if item.evidence_id not in evidence_ids:
                    continue
                if item.source_kind != "external":
                    raise ValueError("Research source has an invalid evidence kind")
                previous = external.get(item.evidence_id)
                if previous is not None and previous != item:
                    raise ValueError(
                        "Research evidence identity changed across snapshots"
                    )
                external[item.evidence_id] = item
        artifact_ids = {item.id for item in evidence}
        if (
            artifact_ids & external.keys()
            or evidence_ids != artifact_ids | external.keys()
        ):
            reject(
                "evidence_unavailable",
                "Export cited evidence is unavailable in the selected version",
                409,
            )
        snapshot: dict[str, Any] = {
            "proposals": frozen_proposals.model_dump(mode="json"),
            "config": deepcopy(run.config),
            "decisions": decisions,
            "source_versions": [
                {
                    "source_version_id": value.id,
                    "source_id": value.source_id,
                    "sha256": value.sha256,
                    "object_key": value.object_key,
                }
                for value in source_versions
            ],
            "evidence": [
                {
                    "evidence_id": value.id,
                    "source_version_id": value.source_version_id,
                    "locator": value.locator,
                    "scope": value.scope,
                    "excerpt": value.excerpt,
                    "details": value.details,
                    "limitations": value.limitations,
                }
                for value in evidence
            ]
            + [value.model_dump(mode="json") for value in external.values()],
            "analysis": analysis,
            "renders": [
                render_snapshot(value.payload, run_id)
                for value in outputs
                if value.stage_key.startswith("render:")
                and value.payload["manifest"]["proposal_set_version"]
                == body.proposal_set_version
                and set(value.payload["manifest"]["variant_ids"]).issubset(
                    body.variant_ids
                )
            ],
            "audit": [
                {
                    "audit_id": value.id,
                    "actor_id": value.actor_id,
                    "actor_type": value.actor_type,
                    "action": value.action,
                    "resource_id": value.resource_id,
                    "rationale": value.rationale,
                    "proposal_set_version": value.proposal_set_version,
                    "occurred_at": value.occurred_at.isoformat(),
                }
                for value in history
            ],
            "cutoff": cutoff.isoformat(),
        }
        manifest: dict[str, str | int | list[str]] = {
            "proposal_set_version": body.proposal_set_version,
            "variant_ids": body.variant_ids,
            "format": body.format,
            "source_version_ids": proposals.source_version_ids,
            "research_ids": proposals.research_ids,
            "scoring_policy_version": proposals.scoring_policy_version,
            "config_version": proposals.config_version,
            "cutoff": cutoff.isoformat(),
        }
        result = w.ExportResult(
            export_id=command_id, status="queued", manifest=manifest
        )
        snapshot["export_manifest"] = {"export_id": command_id, **manifest}
        session.add(
            StageOutput(
                run_id=run.id,
                generation_id=run.generation_id,
                stage_key=f"export:{command_id}",
                payload={
                    "result": result.model_dump(mode="json"),
                    "snapshot": snapshot,
                    "workspace_id": actor.workspace_id,
                    "object_key": f"{actor.workspace_id}/exports/{command_id}.{body.format}",
                },
            )
        )
        session.add(Outbox(command_id=command_id))
        emit(session, run, "export.updated", "export", command_id)
        return accepted(command_id, run_id, "exports")

    return w.Accepted.model_validate(
        command(
            actor,
            "export",
            key,
            {"run_id": run_id, **body.model_dump(mode="json")},
            effect,
            run_id=run_id,
        )
    )


def safe_cell(value: Any) -> Any:
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False, sort_keys=True)
    if isinstance(value, str) and len(value) > 32767:
        raise ValueError(
            "XLSX cell content exceeds its limit; use JSON for the complete snapshot"
        )
    if isinstance(value, str) and (
        value.startswith(("\t", "\r", "\n"))
        or value.lstrip().startswith(("=", "+", "-", "@"))
    ):
        return "'" + value
    return value


def export_bytes(snapshot: dict[str, Any], format: str) -> bytes:
    if format == "json":
        return json.dumps(
            snapshot, ensure_ascii=False, sort_keys=True, allow_nan=False
        ).encode()
    if format != "xlsx":
        raise ValueError("Unsupported export format")
    workbook = Workbook()
    del workbook[workbook.sheetnames[0]]
    frozen_time = datetime.fromisoformat(snapshot["cutoff"]).replace(tzinfo=None)
    workbook.properties.created = frozen_time
    workbook.properties.modified = frozen_time
    for title, rows in (
        ("Requirements", snapshot["proposals"]["changes"]),
        ("Variants", snapshot["proposals"]["variants"]),
        ("Evidence", snapshot["evidence"]),
        ("Decisions", snapshot["decisions"]),
        ("History", snapshot["audit"]),
        ("Sources", snapshot["source_versions"]),
        (
            "Factors",
            [
                {"change_id": change["change_id"], **change["factors"]}
                for change in snapshot["proposals"]["changes"]
            ],
        ),
        ("Research", [snapshot["analysis"].get("research", {})]),
        ("Renders", snapshot.get("renders", [])),
        (
            "Manifest",
            [
                {
                    "cutoff": snapshot["cutoff"],
                    **snapshot.get("export_manifest", {}),
                    **{
                        key: value
                        for key, value in snapshot["proposals"].items()
                        if key not in {"changes", "variants"}
                    },
                }
            ],
        ),
    ):
        sheet = workbook.create_sheet(title)
        headers = list(dict.fromkeys(key for row in rows for key in row))
        if headers:
            sheet.append(headers)
            for row in rows:
                sheet.append([safe_cell(row.get(key)) for key in headers])
    result = BytesIO()
    workbook.save(result)
    workbook.close()
    deterministic = BytesIO()
    with ZipFile(result) as original, ZipFile(deterministic, "w") as frozen:
        for item in original.infolist():
            item.date_time = (1980, 1, 1, 0, 0, 0)
            content = original.read(item.filename)
            if item.filename == "docProps/core.xml":
                core = ElementTree.fromstring(content)
                modified = core.find("{http://purl.org/dc/terms/}modified")
                if modified is not None:
                    modified.text = frozen_time.isoformat(timespec="seconds") + "Z"
                content = ElementTree.tostring(core, encoding="utf-8")
            frozen.writestr(item, content)
    return deterministic.getvalue()


@activity.defn
def export_build(args: list[str]) -> str:
    run_id, generation, export_id = args
    with Session() as session:
        row = resource_row(session, run_id, f"export:{export_id}")
        payload = deepcopy(row.payload)
        result = w.ExportResult.model_validate(payload["result"])
        if result.status == "ready":
            return "ready"
    try:
        data = export_bytes(payload["snapshot"], str(result.manifest["format"]))
        store = storage(Settings())
        try:
            existing = store.read(payload["object_key"])
        except (FileNotFoundError,):
            existing = None
        except Exception as error:
            response = getattr(error, "response", {})
            if response.get("Error", {}).get("Code") not in {"NoSuchKey", "404"}:
                raise
            existing = None
        if existing is None:
            store.put(payload["object_key"], data)
        elif existing != data:
            raise ValueError(
                "Existing export content does not match the immutable snapshot"
            )
        result.status = "ready"
        result.download_url = f"/api/v1/runs/{run_id}/exports/{export_id}/download"
    except Exception:
        result.status = "failed"
        result.download_url = None
    with Session.begin() as session:
        run = session.scalar(select(Run).where(Run.id == run_id).with_for_update())
        row = resource_row(session, run_id, f"export:{export_id}")
        if row.payload["result"]["status"] == "ready":
            return "ready"
        payload["result"] = result.model_dump(mode="json")
        row.payload = payload
        if run is not None:
            emit(session, run, "export.updated", "export", export_id)
    return result.status


def get_export(actor: Actor, run_id: str, export_id: str) -> w.ExportResult:
    with Session() as session:
        owned_run(session, actor, run_id)
        return w.ExportResult.model_validate(
            resource_row(session, run_id, f"export:{export_id}").payload["result"]
        )


def download_export(actor: Actor, run_id: str, export_id: str) -> tuple[bytes, str]:
    with Session.begin() as session:
        owned_run(session, actor, run_id)
        row = resource_row(session, run_id, f"export:{export_id}")
        result = w.ExportResult.model_validate(row.payload["result"])
        if result.status != "ready":
            reject("export_not_ready", "Export is not ready for download", 409)
        key = str(row.payload["object_key"])
        audit(session, actor, "export.download", export_id, run_id)
    data = storage(Settings()).read(key)
    media_type = (
        "application/json"
        if result.manifest["format"] == "json"
        else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    return data, media_type
