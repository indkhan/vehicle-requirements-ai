from copy import deepcopy
from io import BytesIO
from uuid import uuid4

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from app import wire as w
from app.control import apply_instruction
from app.exports import export_bytes
from app.rendering import build_manifest, queued_result, validate_result


def sample():
    change = w.Change(
        change_id="change_a",
        requirement_id="requirement_a",
        version=1,
        origin="ai",
        outcome="Customer comfort",
        part_ids=["part_seat"],
        scope=w.Scope(cohort_ids=["cohort_a", "cohort_b"]),
        horizon=2030,
        before=None,
        after="Evaluate comfort",
        finding_ids=["finding_a"],
        supporting_evidence_ids=[],
        opposing_evidence_ids=[],
        metric=None,
        unit=None,
        target=None,
        target_basis=None,
        conditions=[],
        validation_method="Customer study",
        target_maturity="needs_validation",
        factors=w.Factors(
            impact=3, breadth=2, gap=1, future_relevance=4, rationales={}
        ),
        score=50,
        scoring_policy_version="policy_v1",
        evidence_strength="low",
        uncertainty=[],
    )
    return w.Proposals(
        proposal_set_version=1,
        run_id="run_demo",
        generation_id="generation_a",
        source_version_ids=[],
        research_ids=[],
        config_version=1,
        scoring_policy_version="policy_v1",
        variants=[
            w.Variant(
                variant_id="variant_a",
                title="Initial",
                cohort_ids=["cohort_a", "cohort_b"],
                assumptions=[],
                requirement_ids=[change.requirement_id],
                change_ids=[change.change_id],
                related_variant_ids=[],
            )
        ],
        changes=[change],
        unresolved_conflict_ids=[],
        gaps=[],
    )


@pytest.mark.parametrize(
    "operation,fields,targets,intent",
    [
        ("edit_change", {"after": "PM revision"}, ["change_a"], None),
        ("remove_change", {}, ["change_a"], None),
        ("add_change", {"after": "PM idea"}, ["variant_a"], "Comfort idea"),
        (
            "edit_target",
            {
                "target": 3,
                "metric": "rating",
                "unit": "points",
                "target_basis": "Proposed study",
            },
            ["change_a"],
            None,
        ),
        ("regroup_variants", {"change_ids": ["change_a"]}, ["variant_a"], None),
        ("request_more_variants", {}, ["variant_a"], "Scoped comfort alternative"),
    ],
)
def test_six_pm_operations_preserve_original_and_publish_a_new_version(
    operation, fields, targets, intent
):
    base = sample()
    original = base.model_dump()
    result = apply_instruction(
        base,
        w.Instruction(
            operation=operation,
            base_proposal_set_version=1,
            target_ids=targets,
            changed_fields=fields,
            intent=intent,
        ),
        "generation_b",
        2030,
    )
    assert base.model_dump() == original
    assert result.proposal_set_version == 2
    assert result.generation_id == "generation_b"
    if operation in {"edit_change", "edit_target"}:
        assert result.changes[0].origin == "pm"
        assert result.changes[0].version == 2
        assert result.changes[0].factors.impact is None
        assert result.changes[0].decision_state == "proposed"
    if operation == "remove_change":
        assert result.changes == [] and result.variants[0].change_ids == []
    if operation == "add_change":
        assert result.changes[-1].horizon == 2030
        assert result.changes[-1].origin == "pm"
        assert result.changes[-1].supporting_evidence_ids == []
    if operation == "request_more_variants":
        assert len(result.variants) == 3
        assert {tuple(value.cohort_ids) for value in result.variants[1:]} == {
            ("cohort_a",),
            ("cohort_b",),
        }
        assert result.gaps


def test_pm_cannot_inject_approval_or_unsupported_numeric_target():
    for fields in ({"decision_state": "approved"}, {"target": 3}):
        with pytest.raises(HTTPException):
            apply_instruction(
                sample(),
                w.Instruction(
                    operation="edit_change",
                    base_proposal_set_version=1,
                    target_ids=["change_a"],
                    changed_fields=fields,
                ),
                "generation_b",
            )


def test_frontend_cohort_regroup_preserves_evidence_and_rejects_invented_scopes():
    base = sample()
    instruction = w.Instruction(
        operation="regroup_variants",
        base_proposal_set_version=1,
        target_ids=["variant_a"],
        changed_fields={"cohort_ids": ["cohort_b"]},
    )
    result = apply_instruction(base, instruction, "generation_b")
    assert result.variants[0].cohort_ids == ["cohort_b"]
    assert result.variants[0].change_ids == base.variants[0].change_ids
    assert result.changes == base.changes
    assert base.variants[0].cohort_ids == ["cohort_a", "cohort_b"]
    instruction.changed_fields = {"cohort_ids": ["invented"]}
    with pytest.raises(HTTPException) as error:
        apply_instruction(base, instruction, "generation_b")
    assert error.value.status_code == 422


def test_more_variants_without_distinct_supported_scope_reports_a_gap():
    base = sample()
    base.variants[0].cohort_ids = ["cohort_a"]
    result = apply_instruction(
        base,
        w.Instruction(
            operation="request_more_variants",
            base_proposal_set_version=1,
            target_ids=[],
            changed_fields={},
            intent="Unverified alternative",
        ),
        "generation_b",
    )
    assert len(result.variants) == 1
    assert result.gaps


def test_renderer_rejects_wrong_version_unknown_changes_and_omitted_terminal_results():
    body = w.RenderCommand(
        proposal_set_version=1, variant_ids=["variant_a"], baseline_asset_id="asset_a"
    )
    manifest = build_manifest(
        sample(),
        body,
        "render_a",
        {"baseline_asset_id": "asset_a", "part_catalog_version": "catalog_a"},
    )
    result = queued_result(manifest)
    result.status = "succeeded"
    for mutation in ("wrong_version", "unknown_change", "missing_change"):
        invalid = result.model_copy(deep=True)
        if mutation == "wrong_version":
            invalid.proposal_set_version = 2
        if mutation == "unknown_change":
            invalid.outcomes = [
                w.RenderOutcome(
                    change_id="unknown", outcome="unsupported", asset_ids=[]
                )
            ]
        with pytest.raises(ValueError):
            validate_result(manifest, invalid)


def test_renderer_mixed_outcomes_are_partial_and_do_not_change_proposals():
    base = sample()
    extra = base.changes[0].model_copy(deep=True)
    extra.change_id = "change_b"
    base.changes.append(extra)
    base.variants[0].change_ids.append(extra.change_id)
    manifest = build_manifest(
        base,
        w.RenderCommand(
            proposal_set_version=1,
            variant_ids=["variant_a"],
            baseline_asset_id="asset_a",
        ),
        "render_a",
        {"baseline_asset_id": "asset_a", "part_catalog_version": "catalog_a"},
    )
    result = queued_result(manifest)
    result.status = "succeeded"
    result.outcomes = [
        w.RenderOutcome(change_id="change_a", outcome="annotated", asset_ids=[]),
        w.RenderOutcome(
            change_id="change_b",
            outcome="unsupported",
            reason="Geometry unavailable",
            asset_ids=[],
        ),
    ]
    assert validate_result(manifest, result).status == "partial"
    assert base.changes[0].decision_state == "proposed"


def test_xlsx_preserves_numbers_nulls_factors_and_neutralizes_untrusted_formulas():
    base = sample().model_dump(mode="json")
    base["changes"][0]["after"] = '\t=HYPERLINK("https://example.invalid")'
    base["changes"][0]["target"] = 3.5
    snapshot = {
        "proposals": base,
        "evidence": [],
        "decisions": [],
        "audit": [],
        "source_versions": [],
        "analysis": {},
        "cutoff": "2026-10-10T12:00:00+00:00",
    }
    original = deepcopy(snapshot)
    binary = export_bytes(snapshot, "xlsx")
    assert binary == export_bytes(snapshot, "xlsx")
    book = load_workbook(BytesIO(binary), data_only=False, keep_links=False)
    sheet = book["Requirements"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    assert sheet.cell(2, headers["target"]).value == 3.5
    assert sheet.cell(2, headers["unit"]).value is None
    assert sheet.cell(2, headers["after"]).value.startswith("'")
    assert sheet.cell(2, headers["after"]).data_type == "s"
    factors = book["Factors"]
    factor_headers = {cell.value: cell.column for cell in factors[1]}
    assert factors.cell(2, factor_headers["impact"]).value == 3
    assert snapshot == original
    book.close()


def ready_run(client):
    from test_ingestion_cache import frozen_run

    from app.db import Run, Session, StageOutput

    run_id, generation, source_version, _ = frozen_run(client)
    proposals = sample()
    proposals.run_id = run_id
    proposals.generation_id = generation
    proposals.source_version_ids = [source_version]
    with Session.begin() as session:
        run = session.get(Run, run_id)
        run.analysis_status = "proposals_ready"
        run.current_proposal_set_version = 1
        session.add(
            StageOutput(
                run_id=run_id,
                generation_id=generation,
                stage_key="proposals:1",
                payload=proposals.model_dump(mode="json"),
            )
        )
    return run_id, generation


def test_edit_command_retries_once_and_rejects_concurrent_stale_base(client):
    from app.auth import Actor
    from app.control import get_proposals, instruct, refine_proposals
    from app.db import Run, Session

    actor = Actor("pm_demo", "workspace_demo", "pm")
    run_id, original_generation = ready_run(client)
    body = w.Instruction(
        operation="edit_change",
        base_proposal_set_version=1,
        target_ids=["change_a"],
        changed_fields={"after": "PM revision"},
    )
    key = str(uuid4())
    first = instruct(actor, run_id, body, key)
    assert instruct(actor, run_id, body, key) == first
    with pytest.raises(HTTPException) as error:
        instruct(actor, run_id, body, str(uuid4()))
    assert error.value.status_code == 409
    with Session() as session:
        generation = session.get(Run, run_id).generation_id
    assert generation != original_generation
    assert refine_proposals([run_id, generation, first.command_id]) == "ready"
    assert refine_proposals([run_id, generation, first.command_id]) == "ready"
    assert get_proposals(actor, run_id, 1).changes[0].after == "Evaluate comfort"
    assert get_proposals(actor, run_id).changes[0].after == "PM revision"


def test_export_freezes_decisions_and_download_is_workspace_authorized(client):
    from app.auth import Actor
    from app.control import decide
    from app.exports import create_export, download_export, export_build

    actor = Actor("pm_demo", "workspace_demo", "pm")
    run_id, generation = ready_run(client)
    exported = create_export(
        actor,
        run_id,
        w.ExportCommand(
            proposal_set_version=1, variant_ids=["variant_a"], format="json"
        ),
        str(uuid4()),
    )
    decide(
        actor,
        run_id,
        w.Decision(
            proposal_set_version=1,
            variant_id="variant_a",
            change_ids=["change_a"],
            action="approve",
            rationale="PM checked evidence",
        ),
        str(uuid4()),
    )
    assert export_build([run_id, generation, exported.command_id]) == "ready"
    binary, media = download_export(actor, run_id, exported.resource_id)
    import json

    snapshot = json.loads(binary)
    assert snapshot["decisions"] == []
    assert snapshot["proposals"]["changes"][0]["decision_state"] == "proposed"
    assert media == "application/json"
    with pytest.raises(HTTPException) as error:
        download_export(
            Actor("pm_other", "workspace_other", "pm"), run_id, exported.resource_id
        )
    assert error.value.status_code == 404


@pytest.mark.parametrize("format", ["json", "xlsx"])
def test_export_preserves_external_citations_from_selected_proposal_generation(
    client, format
):
    import json

    from app.auth import Actor
    from app.db import Run, Session, StageOutput
    from app.exports import create_export, download_export, export_build

    actor = Actor("pm_demo", "workspace_demo", "pm")
    run_id, generation = ready_run(client)
    external = w.Evidence(
        evidence_id="external_supported",
        source_version_id="external_version",
        source_kind="external",
        locator=w.WebLocator(
            url="https://example.test/study", passage_start=4, passage_end=23
        ),
        scope=w.Scope(markets=["DE"]),
        excerpt="Original research quote",
        limitations=["Public sample"],
        source_url="https://example.test/study",
    ).model_dump(mode="json")
    opposing = deepcopy(external)
    opposing.update(evidence_id="external_opposing", excerpt="Original counterpoint")
    newer = deepcopy(external)
    newer["excerpt"] = "Newer generation must not replace frozen source"
    with Session.begin() as session:
        row = session.get(StageOutput, (run_id, generation, "proposals:1"))
        proposals = deepcopy(row.payload)
        proposals["changes"][0]["supporting_evidence_ids"] = [external["evidence_id"]]
        proposals["changes"][0]["opposing_evidence_ids"] = [opposing["evidence_id"]]
        row.payload = proposals
        for stage, source in [
            ("research_sources", external),
            ("research:step:1", opposing),
        ]:
            session.add(
                StageOutput(
                    run_id=run_id,
                    generation_id=generation,
                    stage_key=stage,
                    payload={"sources": [{"evidence": source}]},
                )
            )
        current_generation = str(uuid4())
        session.get(Run, run_id).generation_id = current_generation
        session.add(
            StageOutput(
                run_id=run_id,
                generation_id=current_generation,
                stage_key="research_sources",
                payload={"sources": [{"evidence": newer}]},
            )
        )
    job = create_export(
        actor,
        run_id,
        w.ExportCommand(
            proposal_set_version=1, variant_ids=["variant_a"], format=format
        ),
        str(uuid4()),
    )
    assert export_build([run_id, current_generation, job.resource_id]) == "ready"
    binary, _ = download_export(actor, run_id, job.resource_id)
    with Session() as session:
        frozen = session.get(
            StageOutput, (run_id, current_generation, f"export:{job.resource_id}")
        )
        assert {e["evidence_id"] for e in frozen.payload["snapshot"]["evidence"]} == {
            external["evidence_id"],
            opposing["evidence_id"],
        }
    if format == "json":
        evidence = {e["evidence_id"]: e for e in json.loads(binary)["evidence"]}
        assert evidence == {
            external["evidence_id"]: external,
            opposing["evidence_id"]: opposing,
        }
    else:
        book = load_workbook(BytesIO(binary), read_only=True)
        headers, *rows = list(book["Evidence"].values)
        evidence = {
            row[headers.index("evidence_id")]: dict(zip(headers, row, strict=True))
            for row in rows
        }
        for expected in [external, opposing]:
            actual = evidence[expected["evidence_id"]]
            for key, value in expected.items():
                assert (
                    json.loads(actual[key])
                    if isinstance(value, (dict, list))
                    else actual[key]
                ) == value
        book.close()


def test_export_rejects_external_citation_missing_from_selected_generation(client):
    from sqlalchemy import select

    from app.auth import Actor
    from app.db import Session, StageOutput
    from app.exports import create_export

    run_id, generation = ready_run(client)
    with Session.begin() as session:
        row = session.get(StageOutput, (run_id, generation, "proposals:1"))
        proposals = deepcopy(row.payload)
        proposals["changes"][0]["supporting_evidence_ids"] = ["unavailable_external"]
        row.payload = proposals
    with pytest.raises(HTTPException) as error:
        create_export(
            Actor("pm_demo", "workspace_demo", "pm"),
            run_id,
            w.ExportCommand(
                proposal_set_version=1, variant_ids=["variant_a"], format="json"
            ),
            str(uuid4()),
        )
    assert error.value.status_code == 409
    assert error.value.detail["code"] == "evidence_unavailable"
    with Session() as session:
        assert (
            session.scalar(
                select(StageOutput).where(
                    StageOutput.run_id == run_id, StageOutput.stage_key.like("export:%")
                )
            )
            is None
        )


@pytest.mark.parametrize("format", ["json", "xlsx"])
def test_export_includes_baseline_citations_outside_selected_changes(client, format):
    import json

    from sqlalchemy import select
    from test_ingestion_cache import feedback

    from app.auth import Actor
    from app.db import EvidenceArtifact, RunEvidence, RunSource, Session, StageOutput
    from app.exports import create_export, download_export, export_build
    from app.ingestion import publish_batch

    run_id, generation = ready_run(client)
    with Session() as session:
        source = session.scalar(select(RunSource).where(RunSource.run_id == run_id))
        version, binding = source.source_version_id, source.binding
    publish_batch(run_id, generation, version, feedback(binding))
    with Session.begin() as session:
        evidence_id = session.scalar(
            select(RunEvidence.evidence_id).where(RunEvidence.run_id == run_id)
        )
        artifact = session.get(EvidenceArtifact, evidence_id)
        artifact.source_kind = "equipment"
        artifact.locator = w.PDFLocator(page=18).model_dump(mode="json")
        artifact.excerpt = "Comfort seats are optional in the documented market."
        session.add(
            StageOutput(
                run_id=run_id,
                generation_id=generation,
                stage_key="baselines",
                payload={
                    "items": {
                        "finding_a": {
                            "basis_evidence_ids": [evidence_id],
                            "claims": [{"evidence_id": evidence_id}],
                        }
                    }
                },
            )
        )
    actor = Actor("pm_demo", "workspace_demo", "pm")
    job = create_export(
        actor,
        run_id,
        w.ExportCommand(
            proposal_set_version=1, variant_ids=["variant_a"], format=format
        ),
        str(uuid4()),
    )
    assert export_build([run_id, generation, job.resource_id]) == "ready"
    binary, _ = download_export(actor, run_id, job.resource_id)
    if format == "json":
        evidence = json.loads(binary)["evidence"]
    else:
        book = load_workbook(BytesIO(binary), read_only=True)
        exported_rows = list(book["Evidence"].values)
        assert len(exported_rows) == 2
        headers, *rows = exported_rows
        evidence = [dict(zip(headers, row, strict=True)) for row in rows]
        for item in evidence:
            item["locator"] = json.loads(item["locator"])
        book.close()
    assert len(evidence) == 1
    assert evidence[0]["evidence_id"] == evidence_id
    assert evidence[0]["locator"]["page"] == 18
    assert evidence[0]["excerpt"] == artifact.excerpt


def test_export_rejects_missing_analysis_citation(client):
    from app.auth import Actor
    from app.db import Session, StageOutput
    from app.exports import create_export

    run_id, generation = ready_run(client)
    with Session.begin() as session:
        session.add(
            StageOutput(
                run_id=run_id,
                generation_id=generation,
                stage_key="baselines",
                payload={
                    "items": {"finding_a": {"claims": [{"evidence_id": "missing"}]}}
                },
            )
        )
    with pytest.raises(HTTPException) as error:
        create_export(
            Actor("pm_demo", "workspace_demo", "pm"),
            run_id,
            w.ExportCommand(
                proposal_set_version=1, variant_ids=["variant_a"], format="json"
            ),
            str(uuid4()),
        )
    assert error.value.status_code == 409
    assert error.value.detail["code"] == "evidence_unavailable"


def test_unconfigured_renderer_fails_explicitly_retains_proposals_and_retry_identity(
    client, monkeypatch
):
    import app.rendering as rendering
    from app.auth import Actor
    from app.control import get_proposals

    actor = Actor("pm_demo", "workspace_demo", "pm")
    run_id, generation = ready_run(client)
    monkeypatch.setattr(
        rendering,
        "renderer_call",
        lambda *args, **kwargs: (_ for _ in ()).throw(
            ValueError("renderer_unconfigured")
        ),
    )
    request = rendering.create_render(
        actor,
        run_id,
        w.RenderCommand(
            proposal_set_version=1,
            variant_ids=["variant_a"],
            baseline_asset_id="asset_demo_v1",
        ),
        str(uuid4()),
    )
    assert (
        rendering.render_submit([run_id, generation, request.resource_id]) == "failed"
    )
    assert get_proposals(actor, run_id).changes[0].after == "Evaluate comfort"
    retried = rendering.retry_render(actor, run_id, request.resource_id, str(uuid4()))
    assert retried.command_id != request.command_id
    assert retried.resource_id == request.resource_id
